"""Explicit, atomic workbook exchange. SQLite remains authoritative."""

import io
import secrets
import zipfile
import sqlite3
import os
import tempfile
import math
import re
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl.utils.exceptions import InvalidFileException
from datetime import date, time, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import TypeAdapter, EmailStr
from fastapi import HTTPException
from . import database as db
from .security import hash_password

SHEETS = {
    "Users": ("users", "user_id full_name email phone role created_at"),
    "Patients": ("patients", "patient_id user_id profile_details"),
    "Doctors": (
        "doctors",
        "doctor_id user_id specialization qualification experience_years consultation_fee biography",
    ),
    "Availability": (
        "availability",
        "slot_id doctor_id appointment_date start_time end_time status",
    ),
    "Appointments": (
        "appointments",
        "appointment_id patient_id doctor_id slot_id appointment_date appointment_time booking_status created_at",
    ),
    "Payments": (
        "payments",
        "payment_id appointment_id amount payment_method payment_status transaction_reference",
    ),
    "Notifications": (
        "notifications",
        "notification_id user_id appointment_id message notification_status created_at",
    ),
}


def safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "'")):
        return "'" + value
    return value


def export_workbook():
    workbook = Workbook()
    workbook.remove(workbook.active)
    c = db.connect()
    try:
        c.execute("BEGIN")
        for title, (table, cols) in SHEETS.items():
            headers = cols.split()
            sheet = workbook.create_sheet(title)
            sheet.append(headers)
            for row in c.execute(f'SELECT {",".join(headers)} FROM {table}'):
                sheet.append([safe(v) for v in row])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="A88745")
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 24
        output = io.BytesIO()
        workbook.save(output)
    finally:
        c.rollback()
        c.close()
    data = output.getvalue()
    path = db.DB_PATH.parent / "Doctor_Appointment_Database.xlsx"
    fd, tmp = tempfile.mkstemp(dir=path.parent, suffix=".xlsx")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
    return data


def parse(data, filename):
    if (
        not filename
        or not filename.lower().endswith(".xlsx")
        or len(data) > 5 * 1024 * 1024
    ):
        raise ValueError("Upload an .xlsx workbook smaller than 5 MB.")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if (
            len(z.infolist()) > 200
            or sum(i.file_size for i in z.infolist()) > 30 * 1024 * 1024
            or any("vbaproject" in i.filename.lower() for i in z.infolist())
        ):
            raise ValueError("Workbook is too large or contains macros.")
    wb = load_workbook(
        io.BytesIO(data), read_only=True, data_only=False, keep_links=False
    )
    try:
        if set(wb.sheetnames) != set(SHEETS):
            raise ValueError(
                "Workbook must contain exactly the seven exported worksheets."
            )
        parsed = {}
        for title, (_, cols) in SHEETS.items():
            ws = wb[title]
            headers = cols.split()
            if ws.max_row > 10001 or ws.max_column != len(headers):
                raise ValueError(
                    f"{title}: incorrect columns or more than 10,000 records."
                )
            iterator = iter(ws.rows)
            if [c.value for c in next(iterator)] != headers:
                raise ValueError(f"{title}: column headers must match the export.")
            records = []
            ids = set()
            for number, row in enumerate(iterator, 2):
                if all(c.value is None for c in row):
                    continue
                values = []
                for cell in row:
                    if cell.data_type == "f":
                        raise ValueError(
                            f"{title} row {number}: formulas are not allowed."
                        )
                    v = cell.value
                    if (
                        isinstance(v, str)
                        and v.startswith("'")
                        and v[1:].lstrip().startswith(("=", "+", "-", "@", "'"))
                    ):
                        v = v[1:]
                    if v is not None and len(str(v)) > 4000:
                        raise ValueError(f"{title} row {number}: text is too long.")
                    values.append(v)
                record = dict(zip(headers, values))
                key = record[headers[0]]
                if not isinstance(key, str) or not key or len(key) > 100 or key in ids:
                    raise ValueError(f"{title} row {number}: invalid or duplicate ID.")
                ids.add(key)
                for k, v in record.items():
                    if k in ("profile_details",):
                        record[k] = v or ""
                    elif (
                        k == "appointment_id" and title == "Notifications" and v is None
                    ):
                        continue
                    elif v is None:
                        raise ValueError(f"{title} row {number}: {k} is required.")
                    if k in ("appointment_date",):
                        if (
                            not isinstance(v, str)
                            or date.fromisoformat(v).isoformat() != v
                        ):
                            raise ValueError(
                                f"{title} row {number}: dates must use YYYY-MM-DD text."
                            )
                    if k in ("start_time", "end_time", "appointment_time"):
                        if (
                            not isinstance(v, str)
                            or len(v) != 5
                            or time.fromisoformat(v).isoformat(timespec="minutes") != v
                        ):
                            raise ValueError(
                                f"{title} row {number}: times must use HH:MM text."
                            )
                    if k == "created_at":
                        dt = datetime.fromisoformat(v)
                        if dt.tzinfo is None:
                            raise ValueError("Timestamps must include a timezone.")
                if title == "Users":
                    record["email"] = str(
                        TypeAdapter(EmailStr).validate_python(record["email"])
                    ).lower()
                    if not isinstance(record["phone"], str):
                        raise ValueError(
                            f"{title} row {number}: format phone numbers as text."
                        )
                    if (
                        not re.fullmatch(r"[+\d\s()-]{7,20}", record["phone"])
                        or not 2 <= len(str(record["full_name"])) <= 100
                    ):
                        raise ValueError("User name or phone is invalid.")
                for key, value in record.items():
                    if key in ("consultation_fee", "amount"):
                        if (
                            not isinstance(value, (int, float))
                            or not math.isfinite(value)
                            or not 0 <= value <= 100000
                        ):
                            raise ValueError(f"{title}: invalid monetary amount.")
                    elif (
                        key != "experience_years"
                        and value is not None
                        and not isinstance(value, str)
                    ):
                        raise ValueError(f"{title}: {key} must be text.")
                if title == "Doctors" and (
                    not 2 <= len(record["specialization"]) <= 80
                    or not 2 <= len(record["qualification"]) <= 150
                    or not 10 <= len(record["biography"]) <= 3000
                ):
                    raise ValueError("Doctor profile fields have invalid lengths.")
                if title == "Doctors" and (
                    not isinstance(record["experience_years"], int)
                    or isinstance(record["experience_years"], bool)
                ):
                    raise ValueError("Experience must be a whole number.")
                if title == "Payments" and not str(
                    record["transaction_reference"]
                ).startswith("SIM-"):
                    raise ValueError("Only simulated payment references are accepted.")
                records.append(record)
            parsed[title] = records
        return parsed
    finally:
        wb.close()


def import_workbook(data, filename, dry_run=True):
    try:
        parsed = parse(data, filename)
        c = db.connect()
        try:
            c.execute("BEGIN IMMEDIATE")
            for title, (table, cols) in SHEETS.items():
                for record in parsed[title]:
                    record = dict(record)
                    headers = cols.split()
                    existing = c.execute(
                        f"SELECT * FROM {table} WHERE {headers[0]}=?",
                        (record[headers[0]],),
                    ).fetchone()
                    if title == "Users":
                        if existing and existing["role"] != record["role"]:
                            raise ValueError(
                                "Account roles cannot be changed by workbook import."
                            )
                        if record["role"] == "admin":
                            if not existing or any(
                                existing[k] != record[k] for k in headers
                            ):
                                raise ValueError(
                                    "Administrator accounts cannot be created or changed in Excel."
                                )
                        if not existing:
                            record.update(
                                password_hash=hash_password(secrets.token_urlsafe(32)),
                                active=0,
                            )
                    if existing:
                        immutable = {
                            "Patients": ["user_id"],
                            "Doctors": ["user_id"],
                            "Availability": ["doctor_id"],
                            "Appointments": ["patient_id", "doctor_id"],
                            "Payments": ["appointment_id"],
                            "Notifications": ["user_id", "appointment_id"],
                        }.get(title, [])
                        if any(existing[k] != record[k] for k in immutable):
                            raise ValueError(
                                f"{title}: record ownership cannot change."
                            )
                        if (
                            title == "Appointments"
                            and existing["booking_status"] != "confirmed"
                            and any(existing[k] != record[k] for k in headers)
                        ):
                            raise ValueError(
                                "Completed and cancelled appointments are immutable."
                            )
                    if title == "Doctors":
                        c.execute(
                            "INSERT OR IGNORE INTO specializations VALUES(?)",
                            (record["specialization"],),
                        )
                    keys = list(record)
                    if existing:
                        updates = ",".join(f"{k}=?" for k in keys[1:])
                        c.execute(
                            f"UPDATE {table} SET {updates} WHERE {keys[0]}=?",
                            [record[k] for k in keys[1:]] + [record[keys[0]]],
                        )
                    else:
                        c.execute(
                            f'INSERT INTO {table}({",".join(keys)}) VALUES({",".join("?" for _ in keys)})',
                            list(record.values()),
                        )
            checks = [
                (
                    "SELECT 1 FROM patients p JOIN users u USING(user_id) WHERE u.role!='patient'",
                    "Patient accounts must have the patient role.",
                ),
                (
                    "SELECT 1 FROM doctors d JOIN users u USING(user_id) WHERE u.role!='doctor'",
                    "Doctor accounts must have the doctor role.",
                ),
                (
                    "SELECT 1 FROM users u WHERE (role='patient' AND NOT EXISTS(SELECT 1 FROM patients p WHERE p.user_id=u.user_id)) OR (role='doctor' AND NOT EXISTS(SELECT 1 FROM doctors d WHERE d.user_id=u.user_id))",
                    "Every patient/doctor needs a matching profile.",
                ),
                (
                    "SELECT 1 FROM availability a JOIN availability b ON a.doctor_id=b.doctor_id AND a.appointment_date=b.appointment_date AND a.slot_id<b.slot_id AND a.start_time<b.end_time AND a.end_time>b.start_time",
                    "Doctor time slots cannot overlap.",
                ),
                (
                    "SELECT 1 FROM appointments a JOIN availability s USING(slot_id) WHERE a.doctor_id!=s.doctor_id OR a.appointment_date!=s.appointment_date OR a.appointment_time!=s.start_time",
                    "Appointment dates and doctors must match their slots.",
                ),
                (
                    "SELECT 1 FROM availability s WHERE (status='booked') != EXISTS(SELECT 1 FROM appointments a WHERE a.slot_id=s.slot_id AND a.booking_status!='cancelled')",
                    "Slot status must agree with appointment status.",
                ),
                (
                    "SELECT 1 FROM appointments a LEFT JOIN payments p USING(appointment_id) WHERE p.payment_id IS NULL OR (a.booking_status='cancelled' AND p.payment_status!='refunded') OR (a.booking_status!='cancelled' AND p.payment_status!='simulated_paid')",
                    "Every appointment needs a matching simulated payment or refund.",
                ),
            ]
            for sql, message in checks:
                if c.execute(sql).fetchone():
                    raise ValueError(message)
            result = {
                "valid": True,
                "applied": not dry_run,
                "counts": {k: len(v) for k, v in parsed.items()},
                "preview": {k: v[:5] for k, v in parsed.items()},
                "message": "Validation passed. Import merges rows by ID; missing rows are not deleted. New accounts remain inactive until an administrator sets a password.",
            }
            if dry_run:
                c.rollback()
            else:
                c.commit()
            return result
        finally:
            c.rollback()
            c.close()
    except (
        ValueError,
        TypeError,
        sqlite3.IntegrityError,
        zipfile.BadZipFile,
        KeyError,
        StopIteration,
        OverflowError,
        ParseError,
        DefusedXmlException,
        InvalidFileException,
        OSError,
    ) as exc:
        raise HTTPException(422, f"Workbook rejected: {exc}") from exc
