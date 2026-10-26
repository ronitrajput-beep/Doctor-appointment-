"""Patient-maintained history and each doctor's private practice workspace."""

import io
from datetime import timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from fastapi import HTTPException
from . import database as db, appointments as ap
from .excel import safe

HISTORY_FIELDS = (
    "date_of_birth",
    "blood_group",
    "allergies",
    "conditions",
    "medications",
    "surgeries",
    "family_history",
    "other_notes",
)


def doctor_id(user):
    return db.one("SELECT doctor_id FROM doctors WHERE user_id=?", (user["user_id"],))[
        "doctor_id"
    ]


def history(pid):
    result = db.one("SELECT * FROM medical_histories WHERE patient_id=?", (pid,))
    return result or {
        **{f: "" for f in HISTORY_FIELDS},
        "patient_id": pid,
        "updated_at": "",
    }


def history_for_user(user):
    p = db.one("SELECT patient_id FROM patients WHERE user_id=?", (user["user_id"],))
    return history(p["patient_id"])


def save_history(user, body):
    if body.date_of_birth and body.date_of_birth > db.now().date():
        raise HTTPException(422, "Date of birth cannot be in the future.")
    pid = db.one("SELECT patient_id FROM patients WHERE user_id=?", (user["user_id"],))[
        "patient_id"
    ]
    values = body.model_dump(mode="json")
    values["date_of_birth"] = values["date_of_birth"] or ""
    before = history(pid)
    if all(before[field] == values[field] for field in HISTORY_FIELDS):
        return before
    with db.transaction() as c:
        cols = ",".join(HISTORY_FIELDS)
        updates = ",".join(f"{f}=excluded.{f}" for f in HISTORY_FIELDS)
        c.execute(
            f'INSERT INTO medical_histories(patient_id,{cols},updated_at) VALUES({",".join("?" for _ in range(10))}) ON CONFLICT(patient_id) DO UPDATE SET {updates},updated_at=excluded.updated_at',
            (pid, *(values[f] for f in HISTORY_FIELDS), db.stamp()),
        )
        for target in c.execute(
            "SELECT DISTINCT d.user_id FROM appointments a JOIN doctors d USING(doctor_id) JOIN users u ON u.user_id=d.user_id WHERE a.patient_id=? AND a.booking_status!='cancelled' AND u.active=1",
            (pid,),
        ).fetchall():
            db.notify(
                c,
                target[0],
                None,
                f"{user['full_name']} updated their medical history. Review it in My patients.",
            )
    return history(pid)


def patients(user):
    return db.rows(
        """SELECT p.patient_id,u.full_name,u.email,u.phone,
        count(a.appointment_id) AS appointment_count,
        max(a.appointment_date) AS last_appointment,
        coalesce(h.updated_at,'') AS history_updated_at
        FROM appointments a JOIN patients p USING(patient_id) JOIN users u ON u.user_id=p.user_id
        LEFT JOIN medical_histories h USING(patient_id)
        WHERE a.doctor_id=? AND a.booking_status!='cancelled'
        GROUP BY p.patient_id ORDER BY u.full_name""",
        (doctor_id(user),),
    )


def patient_detail(user, pid):
    with db.transaction():
        p = next((p for p in patients(user) if p["patient_id"] == pid), None)
        if not p:
            raise HTTPException(404, "Patient not found in your consultations.")
        return {
            **p,
            "medical_history": history(pid),
            "appointments": [a for a in ap.listing(user) if a["patient_id"] == pid],
        }


def time_off(user):
    return db.rows(
        "SELECT * FROM doctor_time_off WHERE doctor_id=? ORDER BY starts_at DESC",
        (doctor_id(user),),
    )


def add_time_off(user, body):
    starts = ap.instant(
        body.start_date.isoformat(), body.start_time.isoformat(timespec="minutes")
    )
    ends = ap.instant(
        body.end_date.isoformat(), body.end_time.isoformat(timespec="minutes")
    )
    if ends <= starts or ends <= db.now():
        raise HTTPException(
            422,
            "Choose an unavailable period ending in the future and after its start.",
        )
    if ends - starts > timedelta(days=366):
        raise HTTPException(422, "Unavailable periods must be no longer than one year.")
    start, end = starts.strftime("%Y-%m-%dT%H:%M"), ends.strftime("%Y-%m-%dT%H:%M")
    did, lid = doctor_id(user), db.uid("LEAVE")
    with db.transaction() as c:
        if c.execute(
            "SELECT 1 FROM doctor_time_off WHERE doctor_id=? AND starts_at<? AND ends_at>?",
            (did, end, start),
        ).fetchone():
            raise HTTPException(
                409, "This period overlaps an existing unavailable period."
            )
        if c.execute(
            "SELECT 1 FROM availability s JOIN appointments a USING(slot_id) WHERE s.doctor_id=? AND a.booking_status='confirmed' AND s.appointment_date||'T'||s.start_time<? AND s.appointment_date||'T'||s.end_time>?",
            (did, end, start),
        ).fetchone():
            raise HTTPException(
                409,
                "You have confirmed appointments in this period. Arrange changes with the administrator before adding leave.",
            )
        c.execute(
            "INSERT INTO doctor_time_off VALUES(?,?,?,?,?,?)",
            (lid, did, start, end, body.kind, body.note),
        )
    return db.one("SELECT * FROM doctor_time_off WHERE leave_id=?", (lid,))


def remove_time_off(user, lid):
    with db.transaction() as c:
        if not c.execute(
            "DELETE FROM doctor_time_off WHERE leave_id=? AND doctor_id=?",
            (lid, doctor_id(user)),
        ).rowcount:
            raise HTTPException(404, "Unavailable period not found.")
    return {
        "message": "Unavailable period removed. Future unblocked slots can be booked again."
    }


def export(user, kind):
    # One consistent read snapshot. Both identity and row scope come from the session.
    with db.transaction():
        # Helpers use separate read connections; this write lock prevents a concurrent
        # relationship change while the scoped workbook is assembled.
        wb = Workbook()
        wb.remove(wb.active)

        def sheet(name, headers, rows):
            ws = wb.create_sheet(name)
            ws.append(headers)
            for row in rows:
                ws.append([safe(v) for v in row])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="23665E")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 24

        if kind in ("patients", "workspace"):
            rows = []
            for p in patients(user):
                h = history(p["patient_id"])
                rows.append(
                    [
                        p[f]
                        for f in (
                            "patient_id",
                            "full_name",
                            "email",
                            "phone",
                            "appointment_count",
                        )
                    ]
                    + [h[f] for f in HISTORY_FIELDS]
                    + [h["updated_at"]]
                )
            sheet(
                "My Patients",
                [
                    "patient_id",
                    "name",
                    "email",
                    "phone",
                    "appointments",
                    *HISTORY_FIELDS,
                    "history_updated_at",
                ],
                rows,
            )
        if kind in ("appointments", "workspace"):
            headers = [
                "appointment_id",
                "patient_id",
                "patient_name",
                "appointment_date",
                "appointment_time",
                "booking_status",
                "amount",
                "payment_status",
            ]
            sheet(
                "My Appointments",
                headers,
                [[a[f] for f in headers] for a in ap.listing(user)],
            )
        if kind == "workspace":
            headers = [
                "slot_id",
                "appointment_date",
                "start_time",
                "end_time",
                "status",
            ]
            slots = db.rows(
                "SELECT * FROM availability WHERE doctor_id=? ORDER BY appointment_date,start_time",
                (doctor_id(user),),
            )
            sheet("My Availability", headers, [[s[f] for f in headers] for s in slots])
            headers = ["starts_at", "ends_at", "kind", "note"]
            sheet(
                "My Time Off",
                headers,
                [[l[f] for f in headers] for l in time_off(user)],
            )
        output = io.BytesIO()
        wb.save(output)
    return output.getvalue()
