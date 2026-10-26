import io
import sqlite3
from datetime import timedelta
from concurrent.futures import ThreadPoolExecutor
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import database as db, schemas, clinical
from app.main import app
from app.seed import CAST
from test_workflows import slots, book, workbook, upload


def history(client, **values):
    return client.patch("/api/patient/medical-history", json=values)


def period(slot, **values):
    return {
        "start_date": slot["appointment_date"],
        "end_date": slot["appointment_date"],
        "start_time": slot["start_time"],
        "end_time": slot["end_time"],
        "kind": "leave",
        **values,
    }


def test_each_named_doctor_has_distinct_account_and_permissions(env):
    c, login = env
    for i, cast in enumerate(CAST, 1):
        r = c.post(
            "/api/auth/login",
            json={
                "email": "doctor@example.com" if i == 1 else f"doctor{i}@example.com",
                "password": "TestingPass123!",
            },
        )
        assert r.status_code == 200
        assert r.json()["full_name"] == cast["name"] and r.json()["role"] == "doctor"
        assert c.get("/api/doctor/profile").json()["doctor_id"] == f"doctor-{i}"
        assert c.get("/api/admin/users").status_code == 403
        assert book(c, "anything").status_code == 403
    assert (
        len(
            {
                x["user_id"]
                for x in db.rows("SELECT user_id FROM users WHERE role='doctor'")
            }
        )
        == 6
    )


def test_history_saved_without_affecting_profile_and_private_routes(env):
    _, login = env
    p = login()
    before = db.one("SELECT * FROM patients WHERE patient_id='patient-1'")
    assert p.get("/api/patient/medical-history").json()["allergies"] == ""
    r = history(
        p,
        allergies="Penicillin (fictional)",
        conditions="Asthma",
        date_of_birth="2004-05-06",
        blood_group="O+",
    )
    assert r.status_code == 200 and r.json()["updated_at"]
    assert r.json()["date_of_birth"] == "2004-05-06"
    assert p.get("/api/patient/medical-history").json() == r.json()
    assert db.one("SELECT * FROM patients WHERE patient_id='patient-1'") == before
    assert (
        login("patient2").get("/api/patient/medical-history").json()["allergies"] == ""
    )
    for role in ["doctor", "admin"]:
        assert login(role).get("/api/patient/medical-history").status_code == 403
        assert history(login(role), conditions="unauthorized").status_code == 403
    assert history(p, patient_id="patient-2").status_code == 422
    assert history(p, allergies="x" * 3001).status_code == 422
    assert (
        history(
            p, date_of_birth=(db.now().date() + timedelta(days=1)).isoformat()
        ).status_code
        == 422
    )
    assert history(p, blood_group="invalid").status_code == 422
    assert history(p).json()["allergies"] == ""  # patient can clear the history


def test_doctor_history_is_scoped_and_cancelled_only_does_not_grant_access(env):
    _, login = env
    assert (
        history(
            login("patient2"), conditions="Patient two private condition"
        ).status_code
        == 200
    )
    d1, d2 = login("doctor"), login("doctor2")
    assert (
        d1.get("/api/doctor/patients/patient-2").json()["medical_history"]["conditions"]
        == "Patient two private condition"
    )
    assert d2.get("/api/doctor/patients/patient-2").status_code == 404
    c = TestClient(app)
    c.post(
        "/api/auth/login",
        json={"email": "doctor3@example.com", "password": "TestingPass123!"},
    )
    assert c.get("/api/doctor/patients").json() == []
    assert c.get("/api/doctor/patients/patient-1").status_code == 404
    assert login("patient2").get("/api/doctor/patients/patient-1").status_code == 403
    assert (
        d1.patch(
            "/api/patient/medical-history", json={"conditions": "doctor edit"}
        ).status_code
        == 403
    )
    assert "medical_history" not in d2.get("/api/appointments").text


def test_cancellation_revokes_new_patient_history_access_and_export(env):
    _, login = env
    p = login("patient2")
    d = login("doctor2")
    assert history(p, allergies="SECRET TWO").status_code == 200
    slot = slots(p, "doctor-2")[0]
    a = book(p, slot["slot_id"]).json()
    assert d.get("/api/doctor/patients/patient-2").status_code == 200
    assert (
        p.patch("/api/appointments/" + a["appointment_id"] + "/cancel").status_code
        == 200
    )
    assert d.get("/api/doctor/patients/patient-2").status_code == 404
    wb = load_workbook(io.BytesIO(d.get("/api/doctor/export/patients").content))
    assert "SECRET TWO" not in str(list(wb.active.values))


def test_history_change_notifies_only_treating_doctors_and_is_idempotent(env):
    _, login = env
    before = {
        r["user_id"]: r["n"]
        for r in db.rows(
            "SELECT user_id,count(*) n FROM notifications GROUP BY user_id"
        )
    }
    p = login()
    r = history(p, medications="Do not put this medicine in notification text")
    assert r.status_code == 200
    notices = db.rows(
        "SELECT * FROM notifications WHERE message LIKE '%updated their medical history%'"
    )
    assert {n["user_id"] for n in notices} == {"user-doctor-1", "user-doctor-2"}
    assert all("medicine" not in n["message"] for n in notices)
    history(p, medications="Do not put this medicine in notification text")
    assert (
        len(
            db.rows(
                "SELECT * FROM notifications WHERE message LIKE '%updated their medical history%'"
            )
        )
        == 2
    )


def test_doctor_exports_are_scoped_and_formula_safe(env):
    _, login = env
    history(login(), conditions='=HYPERLINK("https://example.test")')
    history(login("patient2"), allergies="OTHER PATIENT PRIVATE")
    d = login("doctor2")
    for kind in ["patients", "appointments", "workspace"]:
        r = d.get("/api/doctor/export/" + kind)
        assert r.status_code == 200
        assert "attachment" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        assert all(
            cell.data_type != "f" for sheet in wb for row in sheet for cell in row
        )
        text = str([list(sheet.values) for sheet in wb])
        assert "OTHER PATIENT PRIVATE" not in text
        assert "password_hash" not in text and "aurelia_session" not in text
        if kind in ("patients", "workspace"):
            rows = list(wb["My Patients"].values)
            assert len(rows) == 2 and rows[1][0] == "patient-1"
            assert rows[1][list(rows[0]).index("conditions")].startswith("'=HYPERLINK")
        if kind == "workspace":
            assert wb.sheetnames == [
                "My Patients",
                "My Appointments",
                "My Availability",
                "My Time Off",
            ]
            assert len(list(wb["My Appointments"].values)) == 2
    for role in ["patient", "admin"]:
        assert login(role).get("/api/doctor/export/workspace").status_code == 403
    assert d.get("/api/doctor/export/database").status_code == 404
    assert d.get("/api/admin/excel/export").status_code == 403


def test_leave_hides_existing_slots_blocks_booking_and_restores_after_removal(env):
    _, login = env
    p, d = login(), login("doctor")
    s = slots(p)[0]
    r = d.post("/api/doctor/time-off", json=period(s))
    assert r.status_code == 201, r.text
    assert s["slot_id"] not in [x["slot_id"] for x in slots(p)]
    assert book(p, s["slot_id"]).status_code == 409
    own = next(
        x
        for x in d.get("/api/doctor/availability").json()
        if x["slot_id"] == s["slot_id"]
    )
    assert own["status"] == "available" and own["time_off"] is True
    assert (
        d.post(
            "/api/doctor/availability",
            json={k: s[k] for k in ["appointment_date", "start_time", "end_time"]},
        ).status_code
        == 409
    )
    assert (
        login("doctor2")
        .delete("/api/doctor/time-off/" + r.json()["leave_id"])
        .status_code
        == 404
    )
    assert d.delete("/api/doctor/time-off/" + r.json()["leave_id"]).status_code == 200
    assert s["slot_id"] in [x["slot_id"] for x in slots(p)]
    assert book(p, s["slot_id"]).status_code == 201


def test_leave_rejects_existing_booking_without_changing_data(env):
    _, login = env
    p, d = login(), login("doctor")
    s = slots(p)[0]
    assert book(p, s["slot_id"]).status_code == 201
    before = db.rows("SELECT * FROM availability ORDER BY rowid")
    r = d.post("/api/doctor/time-off", json=period(s))
    assert r.status_code == 409
    assert "confirmed appointments" in r.json()["detail"]
    assert d.get("/api/doctor/time-off").json() == []
    assert db.rows("SELECT * FROM availability ORDER BY rowid") == before


def test_leave_range_overlap_boundaries_and_scope(env):
    _, login = env
    p, d = login(), login("doctor")
    s = slots(p)[0]
    assert (
        d.post(
            "/api/doctor/time-off", json=period(s, end_time=s["start_time"])
        ).status_code
        == 422
    )
    assert (
        d.post(
            "/api/doctor/time-off",
            json=period(s, start_date="2000-01-01", end_date="2000-01-01"),
        ).status_code
        == 422
    )
    assert (
        d.post(
            "/api/doctor/time-off", json=period(s, start_time="09:00:01")
        ).status_code
        == 422
    )
    r = d.post("/api/doctor/time-off", json=period(s, kind="unavailable"))
    assert r.status_code == 201
    assert d.post("/api/doctor/time-off", json=period(s)).status_code == 409
    adjacent = period(s, start_time=s["end_time"], end_time="10:00")
    assert d.post("/api/doctor/time-off", json=adjacent).status_code == 201
    assert login("doctor2").get("/api/doctor/time-off").json() == []
    assert (
        login("patient").post("/api/doctor/time-off", json=period(s)).status_code == 403
    )
    assert login("patient").get("/api/doctor/time-off").status_code == 403


def test_blocked_slot_does_not_reopen_when_leave_is_removed(env):
    _, login = env
    p, d = login(), login("doctor")
    s = slots(p)[0]
    assert d.delete("/api/doctor/availability/" + s["slot_id"]).status_code == 200
    l = d.post("/api/doctor/time-off", json=period(s))
    assert l.status_code == 201
    assert d.delete("/api/doctor/time-off/" + l.json()["leave_id"]).status_code == 200
    assert s["slot_id"] not in [x["slot_id"] for x in slots(p)]
    payload = {k: s[k] for k in ["appointment_date", "start_time", "end_time"]}
    assert d.post("/api/doctor/availability", json=payload).status_code == 201
    assert s["slot_id"] in [x["slot_id"] for x in slots(p)]


def test_notification_read_all_and_doctor_upcoming_reminders(env):
    _, login = env
    d = login("doctor")
    with db.transaction() as c:
        db.notify(c, "user-doctor-1", None, "Doctor only alert")
        db.notify(c, "user-patient", None, "Patient only alert")
    notices = d.get("/api/notifications").json()
    assert any(n["message"] == "Doctor only alert" for n in notices)
    assert not any(n["message"] == "Patient only alert" for n in notices)
    assert d.patch("/api/notifications/read-all").status_code == 200
    assert not any(
        n["notification_status"] == "unread" for n in d.get("/api/notifications").json()
    )
    assert (
        db.one(
            "SELECT notification_status FROM notifications WHERE message='Patient only alert'"
        )["notification_status"]
        == "unread"
    )


def test_restart_keeps_history_leave_and_credentials(env):
    from app.seed import seed

    _, login = env
    p, d = login(), login("doctor")
    s = slots(p)[0]
    history(p, conditions="Saved condition")
    l = d.post("/api/doctor/time-off", json=period(s))
    assert l.status_code == 201
    before = db.rows("SELECT user_id,password_hash FROM users ORDER BY user_id")
    db.initialize()
    seed()
    db.initialize()
    assert (
        login().get("/api/patient/medical-history").json()["conditions"]
        == "Saved condition"
    )
    assert (
        login("doctor").get("/api/doctor/time-off").json()[0]["leave_id"]
        == l.json()["leave_id"]
    )
    assert db.rows("SELECT user_id,password_hash FROM users ORDER BY user_id") == before


def test_workbook_cannot_import_confirmed_booking_onto_leave(env):
    _, login = env
    p, d, a = login(), login("doctor"), login("admin")
    s = slots(p)[0]
    wb = workbook(a)
    assert d.post("/api/doctor/time-off", json=period(s)).status_code == 201
    wb["Appointments"].append(
        [
            "APT-LEAVE-IMPORT",
            "patient-1",
            "doctor-1",
            s["slot_id"],
            s["appointment_date"],
            s["start_time"],
            "confirmed",
            db.stamp(),
        ]
    )
    wb["Payments"].append(
        [
            "P-LEAVE-IMPORT",
            "APT-LEAVE-IMPORT",
            1200,
            "simulation",
            "simulated_paid",
            "SIM-LEAVE-IMPORT",
        ]
    )
    headers = [c.value for c in wb["Availability"][1]]
    for row in list(wb["Availability"].rows)[1:]:
        if row[0].value == s["slot_id"]:
            row[headers.index("status")].value = "booked"
    before = db.rows("SELECT * FROM appointments ORDER BY rowid")
    r = upload(a, wb, "import")
    assert r.status_code == 422
    assert "unavailable" in r.json()["detail"]
    assert db.rows("SELECT * FROM appointments ORDER BY rowid") == before


def test_old_database_upgrade_keeps_original_records(env):
    _, login = env
    before = {
        t: db.rows(f"SELECT * FROM {t} ORDER BY rowid")
        for t in ["users", "patients", "appointments", "payments", "availability"]
    }
    with db.transaction() as c:
        c.execute("DROP TABLE medical_histories")
        c.execute("DROP TABLE doctor_time_off")
    db.initialize()
    db.initialize()
    assert {t: db.rows(f"SELECT * FROM {t} ORDER BY rowid") for t in before} == before
    assert login().get("/api/patient/medical-history").status_code == 200
    assert login("doctor").get("/api/doctor/time-off").json() == []


def test_local_doctor_login_guide_matches_accounts_and_does_not_reset_passwords(env):
    from app.seed import write_doctor_logins
    from app.security import hash_password

    source = db.DB_PATH.parent.parent / "DEMO_ACCOUNTS.md"
    rows = db.rows(
        "SELECT email,full_name,password_hash FROM users WHERE role='doctor' ORDER BY user_id"
    )
    source.write_text(
        "| Role | Email | Password |\n|---|---|---|\n"
        + "\n".join(f"| doctor | {r['email']} | TestingPass123! |" for r in rows)
    )
    before = db.rows("SELECT user_id,password_hash FROM users ORDER BY user_id")
    write_doctor_logins()
    guide = db.DB_PATH.parent.parent / "DOCTOR_LOGINS.md"
    assert guide.stat().st_mode & 0o777 == 0o600
    text = guide.read_text()
    assert all(r["email"] in text and r["full_name"] in text for r in rows)
    assert text.count("TestingPass123!") == 6
    assert db.rows("SELECT user_id,password_hash FROM users ORDER BY user_id") == before
    with db.transaction() as c:
        c.execute(
            "UPDATE users SET password_hash=? WHERE email='doctor@example.com'",
            (hash_password("ResetPass123!"),),
        )
    write_doctor_logins()
    text = guide.read_text()
    assert (
        "| Dr. Jhatka | doctor@example.com | Ask the administrator to reset this password."
        in text
    )
    assert text.count("TestingPass123!") == 5
