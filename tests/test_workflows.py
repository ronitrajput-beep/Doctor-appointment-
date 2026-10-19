import io
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from openpyxl import load_workbook
from app import database as db
from app.main import app
from fastapi.testclient import TestClient


def slots(client, doctor="doctor-1"):
    r = client.get(f"/api/doctors/{doctor}/availability")
    assert r.status_code == 200
    return r.json()


def book(client, sid, payment="success"):
    return client.post(
        "/api/appointments", json={"slot_id": sid, "payment_result": payment}
    )


def upload(client, wb, action="validate"):
    b = io.BytesIO()
    wb.save(b)
    return client.post(
        "/api/admin/excel/" + action,
        files={
            "file": (
                "data.xlsx",
                b.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def workbook(client):
    r = client.get("/api/admin/excel/export")
    assert r.status_code == 200
    return load_workbook(io.BytesIO(r.content))


def test_authentication_and_password_hash(env):
    client, login = env
    assert client.get("/api/appointments").status_code == 401
    data = {
        "full_name": "Test Patient",
        "email": "new@example.com",
        "phone": "+91 99887 76655",
        "password": "NewPass123!",
    }
    r = client.post("/api/auth/register", json=data)
    assert r.status_code == 201
    assert "httponly" in r.headers["set-cookie"].lower()
    assert "password_hash" not in r.json()
    assert db.one("SELECT password_hash FROM users WHERE email=?", (data["email"],))[
        "password_hash"
    ].startswith("scrypt$")
    assert client.post("/api/auth/register", json=data).status_code == 409
    assert client.get("/api/auth/me").json()["email"] == data["email"]
    token = client.cookies.get("aurelia_session")
    assert client.post("/api/auth/logout").status_code == 200
    client.cookies.set("aurelia_session", token)
    assert client.get("/api/auth/me").status_code == 401


def test_auth_validation_roles_and_origin(env):
    client, login = env
    assert (
        client.post(
            "/api/auth/login",
            json={"email": "patient@example.com", "password": "wrong"},
        ).status_code
        == 401
    )
    assert (
        client.post(
            "/api/auth/register",
            json={
                "full_name": "Bad",
                "email": "x@example.com",
                "phone": "1234567890",
                "password": "Test12345",
                "role": "admin",
            },
        ).status_code
        == 422
    )
    patient = login()
    assert patient.get("/api/admin/users").status_code == 403
    assert patient.post("/api/doctor/availability", json={}).status_code == 403
    assert (
        patient.patch(
            "/api/auth/profile", headers={"origin": "https://malicious.test"}, json={}
        ).status_code
        == 403
    )


def test_booking_and_persistence(env):
    _, login = env
    p = login()
    sid = slots(p)[0]["slot_id"]
    a = book(p, sid)
    assert a.status_code == 201, a.text
    assert a.json()["payment_status"] == "simulated_paid"
    assert (
        db.one("SELECT status FROM availability WHERE slot_id=?", (sid,))["status"]
        == "booked"
    )
    assert book(p, sid).status_code == 409
    db.initialize()
    other = login()
    assert (
        other.get("/api/appointments/" + a.json()["appointment_id"]).status_code == 200
    )
    assert other.get("/api/notifications").json()


def test_payment_failure_and_invalid_slots(env):
    _, login = env
    p = login()
    sid = slots(p)[0]["slot_id"]
    before = db.one("SELECT count(*) AS n FROM appointments")["n"]
    assert book(p, sid, "failure").status_code == 402
    assert db.one("SELECT count(*) AS n FROM appointments")["n"] == before
    assert (
        db.one("SELECT status FROM availability WHERE slot_id=?", (sid,))["status"]
        == "available"
    )
    assert book(p, "missing").status_code == 404
    assert book(login("doctor"), sid).status_code == 403
    with db.transaction() as c:
        c.execute(
            "UPDATE availability SET appointment_date='2000-01-01' WHERE slot_id=?",
            (sid,),
        )
    assert book(p, sid).status_code == 422


def test_cancel_refund_and_rebook(env):
    _, login = env
    p = login()
    sid = slots(p)[0]["slot_id"]
    a = book(p, sid).json()
    r = p.patch("/api/appointments/" + a["appointment_id"] + "/cancel")
    assert r.status_code == 200, r.text
    assert r.json()["payment_status"] == "refunded"
    assert (
        p.patch("/api/appointments/" + a["appointment_id"] + "/cancel").status_code
        == 409
    )
    assert book(login("patient2"), sid).status_code == 201


def test_atomic_reschedule(env):
    _, login = env
    p = login()
    s = slots(p)
    a = book(p, s[0]["slot_id"]).json()
    url = "/api/appointments/" + a["appointment_id"] + "/reschedule"
    assert (
        p.patch(url, json={"slot_id": slots(p, "doctor-2")[0]["slot_id"]}).status_code
        == 422
    )
    assert (
        db.one("SELECT status FROM availability WHERE slot_id=?", (s[0]["slot_id"],))[
            "status"
        ]
        == "booked"
    )
    booked = book(login("patient2"), s[1]["slot_id"])
    assert booked.status_code == 201
    assert p.patch(url, json={"slot_id": s[1]["slot_id"]}).status_code == 409
    r = p.patch(url, json={"slot_id": s[2]["slot_id"]})
    assert r.status_code == 200, r.text
    assert (
        db.one("SELECT status FROM availability WHERE slot_id=?", (s[0]["slot_id"],))[
            "status"
        ]
        == "available"
    )
    assert (
        db.one("SELECT status FROM availability WHERE slot_id=?", (s[2]["slot_id"],))[
            "status"
        ]
        == "booked"
    )


def test_ownership_and_cutoff(env):
    _, login = env
    p = login()
    a = book(p, slots(p)[0]["slot_id"]).json()
    url = "/api/appointments/" + a["appointment_id"]
    assert login("patient2").get(url).status_code == 404
    assert login("doctor2").get(url).status_code == 404
    assert login("doctor").get(url).status_code == 200
    assert login("doctor").patch(url + "/cancel").status_code == 403
    soon = db.now() + timedelta(minutes=50)
    with db.transaction() as c:
        c.execute(
            "UPDATE appointments SET appointment_date=?,appointment_time=? WHERE appointment_id=?",
            (soon.date().isoformat(), soon.strftime("%H:%M"), a["appointment_id"]),
        )
    assert p.patch(url + "/cancel").status_code == 422
    assert (
        p.patch(
            url + "/reschedule", json={"slot_id": slots(p)[1]["slot_id"]}
        ).status_code
        == 422
    )


def test_doctor_schedule_and_completion(env):
    _, login = env
    d = login("doctor")
    p = login()
    date = (db.now().date() + timedelta(days=20)).isoformat()
    data = {"appointment_date": date, "start_time": "10:00", "end_time": "10:30"}
    r = d.post("/api/doctor/availability", json=data)
    assert r.status_code == 201, r.text
    sid = r.json()["slot_id"]
    assert (
        d.post(
            "/api/doctor/availability",
            json={**data, "start_time": "10:15", "end_time": "10:45"},
        ).status_code
        == 409
    )
    assert login("doctor2").delete("/api/doctor/availability/" + sid).status_code == 404
    assert d.delete("/api/doctor/availability/" + sid).status_code == 200
    assert book(p, sid).status_code == 409
    assert d.post("/api/doctor/availability", json=data).status_code == 201
    a = book(p, sid).json()
    url = "/api/appointments/" + a["appointment_id"] + "/complete"
    assert d.delete("/api/doctor/availability/" + sid).status_code == 409
    assert d.patch(url).status_code == 422
    with db.transaction() as c:
        c.execute(
            "UPDATE appointments SET appointment_date='2000-01-01' WHERE appointment_id=?",
            (a["appointment_id"],),
        )
    assert p.patch(url).status_code == 403
    assert d.patch(url).status_code == 200


def test_concurrent_booking_one_winner(env):
    _, login = env
    p = login()
    sid = slots(p)[0]["slot_id"]
    cookies = dict(p.cookies)

    def attempt(_):
        with TestClient(app) as c:
            c.cookies.update(cookies)
            return book(c, sid).status_code

    with ThreadPoolExecutor(max_workers=8) as pool:
        codes = list(pool.map(attempt, range(8)))
    assert codes.count(201) == 1, codes
    assert codes.count(409) == 7, codes


def test_excel_roundtrip_and_booking_sync(env):
    _, login = env
    admin = login("admin")
    p = login()
    a = book(p, slots(p)[0]["slot_id"]).json()
    wb = workbook(admin)
    assert len(wb.sheetnames) == 7
    assert a["appointment_id"] in [
        row[0].value for row in list(wb["Appointments"].rows)[1:]
    ]
    assert not any(
        "password" in str(c.value).lower() for ws in wb for row in ws for c in row
    )
    assert upload(admin, wb).status_code == 200
    assert upload(admin, wb, "import").status_code == 200


def test_excel_dryrun_and_update(env):
    _, login = env
    admin = login("admin")
    wb = workbook(admin)
    wb["Users"]["B2"] = "Changed Name"
    r = upload(admin, wb)
    assert r.status_code == 200, r.text
    assert (
        db.one("SELECT full_name FROM users WHERE user_id='user-patient'")["full_name"]
        == "Aarav Shah"
    )
    r = upload(admin, wb, "import")
    assert r.status_code == 200, r.text
    assert (
        db.one("SELECT full_name FROM users WHERE user_id='user-patient'")["full_name"]
        == "Changed Name"
    )


def test_excel_new_account_is_inactive(env):
    _, login = env
    admin = login("admin")
    wb = workbook(admin)
    wb["Users"].append(
        [
            "new-user",
            "New Patient",
            "imported@example.com",
            "+91 99887 76655",
            "patient",
            db.stamp(),
        ]
    )
    wb["Patients"].append(["new-patient", "new-user", "Sample notes"])
    r = upload(admin, wb, "import")
    assert r.status_code == 200, r.text
    assert "password_hash" not in str(r.json())
    assert db.one("SELECT active FROM users WHERE user_id='new-user'")["active"] == 0
    assert (
        admin.patch(
            "/api/admin/users/new-user/password", json={"password": "TestingPass123!"}
        ).status_code
        == 200
    )
    c = TestClient(app)
    assert (
        c.post(
            "/api/auth/login",
            json={"email": "imported@example.com", "password": "TestingPass123!"},
        ).status_code
        == 200
    )


def test_excel_invalid_import_is_atomic(env):
    _, login = env
    admin = login("admin")
    for kind in ["foreign", "duplicate", "formula", "role", "overlap", "payment"]:
        wb = workbook(admin)
        wb["Users"]["B2"] = "Must Roll Back"
        if kind == "foreign":
            wb["Patients"]["B2"] = "missing-user"
        if kind == "duplicate":
            wb["Patients"].append([c.value for c in wb["Patients"][2]])
        if kind == "formula":
            wb["Users"]["B2"] = "=1+1"
        if kind == "role":
            wb["Users"]["E2"] = "admin"
        if kind == "overlap":
            r = [c.value for c in wb["Availability"][2]]
            r[0] = "overlap"
            r[3] = "12:15"
            r[4] = "12:45"
            wb["Availability"].append(r)
        if kind == "payment":
            wb["Payments"]["E2"] = "refunded"
        result = upload(admin, wb, "import")
        assert result.status_code == 422, (kind, result.text)
        assert (
            db.one("SELECT full_name FROM users WHERE user_id='user-patient'")[
                "full_name"
            ]
            == "Aarav Shah"
        )


def test_excel_file_limits_and_access(env):
    _, login = env
    admin = login("admin")
    assert login().get("/api/admin/excel/export").status_code == 403
    for filename, data in [
        ("bad.txt", b"hello"),
        ("broken.xlsx", b"hello"),
        ("huge.xlsx", b"a" * (5 * 1024 * 1024 + 1)),
    ]:
        assert (
            admin.post(
                "/api/admin/excel/import", files={"file": (filename, data)}
            ).status_code
            == 422
        )


def test_excel_formula_escape(env):
    _, login = env
    admin = login("admin")
    with db.transaction() as c:
        c.execute(
            "UPDATE patients SET profile_details='=HYPERLINK(\"bad\")' WHERE patient_id='patient-1'"
        )
    wb = workbook(admin)
    assert wb["Patients"]["C2"].data_type != "f"
    assert wb["Patients"]["C2"].value.startswith("'")
    assert upload(admin, wb, "import").status_code == 200
    assert db.one("SELECT profile_details FROM patients WHERE patient_id='patient-1'")[
        "profile_details"
    ].startswith("=")


def test_admin_management_and_revocation(env):
    _, login = env
    admin = login("admin")
    p = login()
    assert admin.get("/api/admin/dashboard").json()["counts"]["doctors"] == 6
    assert (
        admin.patch(
            "/api/admin/users/user-patient/active", json={"active": False}
        ).status_code
        == 200
    )
    assert p.get("/api/auth/me").status_code == 401
    assert (
        admin.patch(
            "/api/admin/users/user-admin/active", json={"active": False}
        ).status_code
        == 422
    )
    assert (
        admin.post(
            "/api/admin/specializations", json={"name": "Ophthalmology"}
        ).status_code
        == 201
    )
    data = {
        "full_name": "Test Doctor",
        "email": "testdoctor@example.com",
        "phone": "+91 99887 76655",
        "password": "TestingPass123!",
        "specialization": "Ophthalmology",
        "qualification": "MBBS, MD",
        "experience_years": 5,
        "consultation_fee": 850,
        "biography": "Fictional specialist for automated tests.",
    }
    r = admin.post("/api/admin/doctors", json=data)
    assert r.status_code == 201, r.text
    assert admin.delete("/api/admin/specializations/Ophthalmology").status_code == 409
    editable = {k: v for k, v in data.items() if k not in ("email", "password")}
    editable["consultation_fee"] = 900
    assert (
        admin.patch(
            "/api/admin/doctors/" + r.json()["doctor_id"], json=editable
        ).status_code
        == 200
    )


def test_profile_notifications_search(env):
    _, login = env
    p = login()
    r = p.patch(
        "/api/auth/profile",
        json={
            "full_name": "Aarav Updated",
            "phone": "+91 98765 43210",
            "profile_details": "Fictional notes",
        },
    )
    assert r.status_code == 200
    assert p.get("/api/auth/me").json()["profile_details"] == "Fictional notes"
    notices = p.get("/api/notifications").json()
    assert notices
    nid = notices[0]["notification_id"]
    assert p.patch("/api/notifications/" + nid + "/read").status_code == 200
    assert (
        login("patient2").patch("/api/notifications/" + nid + "/read").status_code
        == 404
    )
    doctors = p.get("/api/doctors?sort=fee").json()
    assert [d["consultation_fee"] for d in doctors] == sorted(
        d["consultation_fee"] for d in doctors
    )
    assert (
        len(p.get("/api/doctors?specialization=General%20Medicine&q=Jhatka").json())
        == 1
    )
    assert p.get("/api/health").json()["database"] == 1


def test_seed_is_repeatable(env):
    from app.seed import seed, demo_slots

    before = {
        t: db.one(f"SELECT count(*) AS n FROM {t}")["n"]
        for t in ["users", "availability", "appointments"]
    }
    seed()
    demo_slots()
    demo_slots()
    after = {t: db.one(f"SELECT count(*) AS n FROM {t}")["n"] for t in before}
    assert before == after


def test_legacy_cast_upgrade_preserves_bookings_accounts_and_later_edits(env):
    from app.seed import seed, CAST, LEGACY_NAMES, SPECIALTIES

    _, login = env
    patient = login()
    booked = book(patient, slots(patient)[0]["slot_id"]).json()
    with db.transaction() as c:
        for i, name in enumerate(LEGACY_NAMES, 1):
            c.execute(
                "UPDATE users SET full_name=? WHERE user_id=?",
                (name, f"user-doctor-{i}"),
            )
            c.execute(
                "UPDATE doctors SET qualification='MBBS, MD',biography='Original demo biography' WHERE doctor_id=?",
                (f"doctor-{i}",),
            )
    protected = [
        "appointments",
        "payments",
        "patients",
        "availability",
        "notifications",
        "sessions",
    ]
    before = {
        table: db.rows(f"SELECT * FROM {table} ORDER BY rowid") for table in protected
    }
    credentials = db.rows(
        "SELECT user_id,email,password_hash,active FROM users ORDER BY user_id"
    )
    seed()
    for i, (character, specialty) in enumerate(zip(CAST, SPECIALTIES), 1):
        record = db.one(
            "SELECT d.*,u.full_name FROM doctors d JOIN users u USING(user_id) WHERE doctor_id=?",
            (f"doctor-{i}",),
        )
        assert record["full_name"] == character["name"]
        assert record["biography"] == character["biography"]
        assert record["qualification"] == character["qualification"]
        assert record["specialization"] == specialty
        assert record["image"] == i - 1
    assert before == {
        table: db.rows(f"SELECT * FROM {table} ORDER BY rowid") for table in protected
    }
    assert credentials == db.rows(
        "SELECT user_id,email,password_hash,active FROM users ORDER BY user_id"
    )
    assert (
        patient.get("/api/appointments/" + booked["appointment_id"]).json()[
            "doctor_name"
        ]
        == "Jhatka"
    )
    assert login("doctor").get("/api/auth/me").json()["full_name"] == "Jhatka"
    with db.transaction() as c:
        c.execute(
            "UPDATE doctors SET biography='An administrator wrote this bio.',consultation_fee=777 WHERE doctor_id='doctor-1'"
        )
    seed()
    record = db.one(
        "SELECT biography,consultation_fee FROM doctors WHERE doctor_id='doctor-1'"
    )
    assert record == {
        "biography": "An administrator wrote this bio.",
        "consultation_fee": 777,
    }


def test_cast_upgrade_skips_customized_users_and_disabled_demo_mode(env, monkeypatch):
    from app import seed as seeder

    with db.transaction() as c:
        c.execute(
            "UPDATE users SET full_name='My own doctor' WHERE user_id='user-doctor-1'"
        )
        c.execute(
            "UPDATE users SET full_name=? WHERE user_id='user-doctor-2'",
            (seeder.LEGACY_NAMES[1],),
        )
    monkeypatch.setenv("SEED_DEMO", "0")
    seeder.seed()
    assert (
        db.one("SELECT full_name FROM users WHERE user_id='user-doctor-2'")["full_name"]
        == seeder.LEGACY_NAMES[1]
    )
    monkeypatch.setenv("SEED_DEMO", "1")
    monkeypatch.setattr(seeder, "PRODUCTION", True)
    seeder.seed()
    assert (
        db.one("SELECT full_name FROM users WHERE user_id='user-doctor-2'")["full_name"]
        == seeder.LEGACY_NAMES[1]
    )
    monkeypatch.setattr(seeder, "PRODUCTION", False)
    seeder.seed()
    assert (
        db.one("SELECT full_name FROM users WHERE user_id='user-doctor-1'")["full_name"]
        == "My own doctor"
    )
    assert (
        db.one("SELECT full_name FROM users WHERE user_id='user-doctor-2'")["full_name"]
        == "Jagdish Asthana"
    )
